import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generar_excel_asistencia(registros):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control_Asistencia"

    headers = [
        "PÁGINA", "FECHA", "TURNO", "PROYECTO", "SECTOR", "FRENTE", 
        "JEFE DE GRUPO", "DNI JEFE", "TRABAJADOR", "DNI", 
        "HH", "FASE", "CONTROL ACTIVIDAD", "HORA INGRESO", "HORA SALIDA", 
        "H.N.", "HORAS 60%", "HORAS 100%", "TOTAL HORAS", "ESTADO", "OBSERVACIONES"
    ]
    
    ws.append(headers)

    for r in registros:
        ws.append([
            getattr(r, 'pagina', 1),
            r.fecha,
            r.turno,
            r.proyecto,
            r.sector,
            r.frente,
            r.jefe_grupo,
            r.dni_jefe,
            r.nombres_apellidos,
            r.dni,
            r.hh,
            r.fase,
            r.control_actividad,
            r.hora_ingreso,
            r.hora_salida,
            r.hn,
            r.horas_60,
            r.horas_100,
            r.total_horas,
            r.estado,
            r.observaciones
        ])

    # Estilos visuales para cabecera
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()